from __future__ import annotations

import io
import re
import tempfile
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urlparse

import httpx

# Domain patterns -> source tier
# T1: Red Hat official docs/KB/errata
# T2: Kubernetes upstream, OCP GitHub, CNCF
# T3: Partner-specific (set manually or by partner domain config)
# T4: General / unknown
_TIER_RULES: list[tuple[int, list[str]]] = [
    (1, [
        "access.redhat.com", "docs.redhat.com", "docs.openshift.com",
        "redhat.com/en/blog", "connect.redhat.com",
        "catalog.redhat.com", "errata.devel.redhat.com",
    ]),
    (2, [
        "kubernetes.io", "github.com/kubernetes", "github.com/openshift",
        "github.com/operator-framework", "operatorhub.io",
        "pkg.go.dev/k8s.io", "etcd.io", "helm.sh",
        "github.com/k8snetworkplumbingwg", "github.com/metallb",
    ]),
]


def detect_tier(url: str) -> int:
    parsed = urlparse(url)
    full = parsed.netloc + parsed.path
    for tier, patterns in _TIER_RULES:
        for pat in patterns:
            if pat in full:
                return tier
    return 4


class _TextExtractor(HTMLParser):
    def __init__(self):
        super().__init__()
        self._text: list[str] = []
        self._title = ""
        self._in_title = False
        self._skip = False
        self._skip_tags = {"script", "style", "nav", "footer", "header", "noscript"}

    def handle_starttag(self, tag, attrs):
        if tag == "title":
            self._in_title = True
        if tag in self._skip_tags:
            self._skip = True

    def handle_endtag(self, tag):
        if tag == "title":
            self._in_title = False
        if tag in self._skip_tags:
            self._skip = False
        if tag in ("p", "div", "li", "h1", "h2", "h3", "h4", "br", "tr"):
            self._text.append("\n")

    def handle_data(self, data):
        if self._in_title and not self._title:
            self._title = data.strip()
        if not self._skip:
            self._text.append(data)

    def get_text(self) -> str:
        raw = "".join(self._text)
        lines = [line.strip() for line in raw.splitlines()]
        return re.sub(r"\n{3,}", "\n\n", "\n".join(line for line in lines if line))

    def get_title(self) -> str:
        return self._title


def _build_auth_headers(auth_source: dict | None) -> dict:
    if not auth_source:
        return {}
    auth_type = auth_source.get("auth_type", "")
    auth_value = auth_source.get("auth_value", "")
    if not auth_value:
        return {}
    if auth_type == "cookie":
        return {"Cookie": auth_value}
    if auth_type == "bearer":
        return {"Authorization": f"Bearer {auth_value}"}
    if auth_type == "basic":
        return {"Authorization": f"Basic {auth_value}"}
    return {}


def _extract_with_docling(data: bytes, filename: str, max_chars: int) -> tuple[str, str] | None:
    """Try docling for structured markdown extraction. Returns None if unavailable."""
    try:
        from docling.document_converter import DocumentConverter
    except ImportError:
        return None
    suffix = Path(filename).suffix or ".pdf"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(data)
        tmp_path = tmp.name
    try:
        result = DocumentConverter().convert(tmp_path)
        md = result.document.export_to_markdown()
        title = result.document.name or ""
        return md[:max_chars], title
    except Exception:
        return None
    finally:
        Path(tmp_path).unlink(missing_ok=True)


def _extract_pdf(data: bytes, max_chars: int) -> tuple[str, str]:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    title = reader.metadata.title if reader.metadata and reader.metadata.title else ""
    pages = []
    for page in reader.pages:
        t = page.extract_text()
        if t:
            pages.append(t)
        if sum(len(p) for p in pages) >= max_chars:
            break
    return "\n\n".join(pages)[:max_chars], title


def _extract_docx(data: bytes, max_chars: int) -> tuple[str, str]:
    from docx import Document
    doc = Document(io.BytesIO(data))
    title = doc.core_properties.title or ""
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
        if sum(len(p) for p in parts) >= max_chars:
            break
    return "\n\n".join(parts)[:max_chars], title


def _extract_pptx(data: bytes, max_chars: int) -> tuple[str, str]:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    title = ""
    parts = []
    for i, slide in enumerate(prs.slides):
        slide_text = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        slide_text.append(t)
        if slide_text:
            if i == 0 and not title:
                title = slide_text[0]
            parts.append(f"[Slide {i+1}]\n" + "\n".join(slide_text))
        if sum(len(p) for p in parts) >= max_chars:
            break
    return "\n\n".join(parts)[:max_chars], title


def _extract_xlsx(data: bytes, max_chars: int) -> tuple[str, str]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                rows.append(" | ".join(cells))
            if sum(len(r) for r in rows) >= max_chars:
                break
        if rows:
            parts.append(f"[{sheet}]\n" + "\n".join(rows))
        if sum(len(p) for p in parts) >= max_chars:
            break
    wb.close()
    return "\n\n".join(parts)[:max_chars], wb.sheetnames[0] if wb.sheetnames else ""


def fetch_and_extract(url: str, max_chars: int = 16000,
                      auth_source: dict | None = None) -> dict:
    """Fetch a URL and extract text content, title, and auto-detect tier."""
    tier = detect_tier(url)
    parsed = urlparse(url)
    headers = {"User-Agent": "Mozilla/5.0 (compatible; OCPArchitect/1.0)"}
    if "sharepoint.com" in parsed.netloc:
        headers["Origin"] = f"{parsed.scheme}://{parsed.netloc}"
        headers["Referer"] = f"{parsed.scheme}://{parsed.netloc}/"
    headers.update(_build_auth_headers(auth_source))

    try:
        resp = httpx.get(url, follow_redirects=True, timeout=60, headers=headers)
        resp.raise_for_status()
    except Exception as e:
        return {"title": "", "content": "", "tier": tier, "error": str(e)}

    content_type = resp.headers.get("content-type", "")
    title = ""

    if "text/html" in content_type:
        parser = _TextExtractor()
        parser.feed(resp.text)
        text = parser.get_text()[:max_chars]
        title = parser.get_title()
    elif "text/plain" in content_type or "application/json" in content_type:
        text = resp.text[:max_chars]
    elif "application/pdf" in content_type:
        result = _extract_with_docling(resp.content, url.rsplit("/", 1)[-1], max_chars)
        if result:
            text, title = result
        else:
            text, title = _extract_pdf(resp.content, max_chars)
    elif "wordprocessingml" in content_type or url.lower().endswith(".docx"):
        result = _extract_with_docling(resp.content, url.rsplit("/", 1)[-1], max_chars)
        if result:
            text, title = result
        else:
            text, title = _extract_docx(resp.content, max_chars)
    elif "presentationml" in content_type or url.lower().endswith(".pptx"):
        result = _extract_with_docling(resp.content, url.rsplit("/", 1)[-1], max_chars)
        if result:
            text, title = result
        else:
            text, title = _extract_pptx(resp.content, max_chars)
    elif "spreadsheetml" in content_type or url.lower().endswith(".xlsx"):
        result = _extract_with_docling(resp.content, url.rsplit("/", 1)[-1], max_chars)
        if result:
            text, title = result
        else:
            text, title = _extract_xlsx(resp.content, max_chars)
    else:
        text = ""

    if not title:
        title = url.rsplit("/", 1)[-1].split("?")[0]

    return {"title": title, "content": text, "tier": tier, "error": ""}
